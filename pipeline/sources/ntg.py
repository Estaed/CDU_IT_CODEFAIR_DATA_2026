"""NT Government source: the 2019, 2021 and 2022 mobile coverage lists and the small-cell list.

Reads the four frozen spreadsheets written by ``pipeline.fetch.ntg`` and emits the *listed*
publisher line, keyed on ``bushtel_id``: whether each community appears on each list (matched
by normalised name and BushTel aliases), the 2022 macro / small cell / proximity basis and
provider, and the 2019 backhaul type. Ported from ``spike/lane_rrl_ntg.py`` part (b)
(2026-09-12) with the matching rules unchanged.

Source: data.nt.gov.au (Northern Territory Government open data, CC BY), resolved through the
CKAN ``package_show`` API. Packages:
    2019       list-of-remote-communities-with-mobile-coverage
    2021       remote-communities-with-mobile-coverage
    2022       mobile-phone-coverage-in-remote-areas-of-the-nt
    smallcell  remote-sites-with-mobile-phone-small-cell-coverage
Fetch date 2026-09-12. Column layouts and header offsets:
reports/2026-09-12-arena-reconcile.md S4.
"""

from __future__ import annotations

import math
import re
from pathlib import Path

import openpyxl
import pandas as pd

FETCH_COMMAND = "PYTHONUTF8=1 .venv/Scripts/python -m pipeline.fetch.ntg"
EARTH_RADIUS_KM = 6371.0088

# List spelling -> BushTel spelling, applied to a list row's name before normalisation.
# The 2026-09-12 run matched 60 (2022) / 57 (2021 with mobile) / 46 (2019) / 2 (small cell)
# rows with normalisation and BushTel aliases alone, so this dict is empty by measurement, not
# by omission. A future alias is one line here, with a comment naming the two spellings.
MANUAL_ALIASES: dict[str, str] = {}

COLUMNS = [
    "bushtel_id",
    "ntg2022_listed",
    "ntg2022_matched_name",
    "ntg2022_site_type",
    "ntg2022_macro",
    "ntg2022_small",
    "ntg2022_proximity",
    "ntg2022_provider",
    "ntg2022_population",
    "ntg2022_coord_dist_km",
    "ntg2021_with_mobile_listed",
    "ntg2021_mobile_phone",
    "ntg2021_comment",
    "ntg2019_listed",
    "ntg2019_backhaul",
    "ntg2019_provider",
    "ntg2019_coord_dist_km",
    "smallcell_listed",
    "smallcell_provider",
]


def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(phi1) * math.cos(phi2) * math.sin(dlambda / 2) ** 2
    return 2 * EARTH_RADIUS_KM * math.asin(math.sqrt(a))


def normalise_name(name: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", str(name).upper())


def _workbook(path: Path) -> openpyxl.Workbook:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(
            f"NT Government coverage list not found: {path}. Re-fetch it with: {FETCH_COMMAND}"
        )
    with path.open("rb") as f:
        return openpyxl.load_workbook(f, data_only=True)


def _rows(ws, first_row: int) -> list[list]:
    """Cell values of every row from ``first_row`` whose first cell is not empty."""
    rows = []
    for r in range(first_row, ws.max_row + 1):
        vals = [c.value for c in ws[r]]
        if not vals or vals[0] is None:
            continue
        rows.append(vals)
    return rows


def load_2019(path: Path) -> list[dict]:
    ws = _workbook(path)["Sheet1"]
    return [
        {"name": v[0], "lat": v[1], "lon": v[2], "backhaul": v[3], "provider": v[4]}
        for v in _rows(ws, 3)
    ]


def load_2021(path: Path) -> tuple[list[dict], list[dict]]:
    wb = _workbook(path)
    with_mobile = [
        {"name": v[0], "type": v[1], "lon": v[2], "lat": v[3]}
        for v in _rows(wb["Communities with Mobile"], 2)
    ]
    gazetteer = [
        {
            "name": v[0],
            "type": v[1],
            "lon": v[2],
            "lat": v[3],
            "mobile_phone": v[4],
            "comment": v[5],
        }
        for v in _rows(wb["Communities"], 2)
    ]
    return with_mobile, gazetteer


def load_2022(path: Path) -> list[dict]:
    ws = _workbook(path)["Communities with Mobile"]
    return [
        {
            "name": v[0],
            "site_type": v[1],
            "population": v[2],
            "macro": v[3],
            "small": v[4],
            "proximity": v[5],
            "provider": v[6],
            "lat": v[7],
            "lon": v[8],
        }
        for v in _rows(ws, 6)
    ]


def load_smallcell(path: Path) -> list[dict]:
    ws = _workbook(path)["Sheet1"]
    return [{"name": v[0], "lat": v[1], "lon": v[2], "provider": v[3]} for v in _rows(ws, 3)]


def match_community(names_norm: set[str], list_rows: list[dict]) -> dict | None:
    """The first list row whose normalised name is one of the community's names, or None."""
    for row in list_rows:
        name = MANUAL_ALIASES.get(row["name"], row["name"])
        if normalise_name(name) in names_norm:
            return row
    return None


def _distance(community: dict, match: dict) -> str:
    if match.get("lat") is None or match.get("lon") is None:
        return ""
    km = haversine_km(community["lat"], community["lon"], float(match["lat"]), float(match["lon"]))
    return str(round(km, 3))


def _row(community: dict, lists: dict[str, list[dict]]) -> dict:
    aliases = [a.strip() for a in str(community["aliases"] or "").split(",") if a.strip()]
    names_norm = {normalise_name(community["name"])} | {normalise_name(a) for a in aliases}
    row: dict = {"bushtel_id": int(community["bushtel_id"])}

    m2022 = match_community(names_norm, lists["2022"])
    if m2022 is not None:
        row.update(
            {
                "ntg2022_listed": "1",
                "ntg2022_matched_name": str(m2022["name"]),
                "ntg2022_site_type": str(m2022["site_type"] or ""),
                "ntg2022_macro": "1" if m2022["macro"] else "0",
                "ntg2022_small": "1" if m2022["small"] else "0",
                "ntg2022_proximity": "1" if m2022["proximity"] else "0",
                "ntg2022_provider": str(m2022["provider"] or ""),
                "ntg2022_population": (
                    str(m2022["population"]) if m2022["population"] is not None else ""
                ),
                "ntg2022_coord_dist_km": _distance(community, m2022),
            }
        )
    else:
        row["ntg2022_listed"] = "0"
        for column in COLUMNS[2:10]:
            row[column] = ""

    m2021_wm = match_community(names_norm, lists["2021_with_mobile"])
    row["ntg2021_with_mobile_listed"] = "1" if m2021_wm is not None else "0"

    m2021_gaz = match_community(names_norm, lists["2021_gazetteer"])
    if m2021_gaz is not None:
        row["ntg2021_mobile_phone"] = str(m2021_gaz["mobile_phone"] or "")
        row["ntg2021_comment"] = str(m2021_gaz["comment"] or "")
    else:
        row["ntg2021_mobile_phone"] = ""
        row["ntg2021_comment"] = ""

    m2019 = match_community(names_norm, lists["2019"])
    if m2019 is not None:
        row["ntg2019_listed"] = "1"
        row["ntg2019_backhaul"] = str(m2019["backhaul"] or "")
        row["ntg2019_provider"] = str(m2019["provider"] or "")
        row["ntg2019_coord_dist_km"] = _distance(community, m2019)
    else:
        row["ntg2019_listed"] = "0"
        row["ntg2019_backhaul"] = ""
        row["ntg2019_provider"] = ""
        row["ntg2019_coord_dist_km"] = ""

    msmall = match_community(names_norm, lists["smallcell"])
    if msmall is not None:
        row["smallcell_listed"] = "1"
        row["smallcell_provider"] = str(msmall["provider"] or "")
    else:
        row["smallcell_listed"] = "0"
        row["smallcell_provider"] = ""
    return row


def load(
    xlsx_2019: Path,
    xlsx_2021: Path,
    xlsx_2022: Path,
    xlsx_smallcell: Path,
    communities: pd.DataFrame,
) -> pd.DataFrame:
    """One row per community in ``communities``, sorted by ``bushtel_id``."""
    with_mobile_2021, gazetteer_2021 = load_2021(xlsx_2021)
    lists = {
        "2019": load_2019(xlsx_2019),
        "2021_with_mobile": with_mobile_2021,
        "2021_gazetteer": gazetteer_2021,
        "2022": load_2022(xlsx_2022),
        "smallcell": load_smallcell(xlsx_smallcell),
    }
    records = communities[["bushtel_id", "name", "aliases", "lat", "lon"]].to_dict("records")
    rows = [_row(community, lists) for community in records]
    rows.sort(key=lambda r: r["bushtel_id"])
    return pd.DataFrame(rows, columns=COLUMNS)
